import re, sys
sys.path.insert(0, "pylib")
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
BASE = "/springbrook/share/eng/esrpxk/results/paper/"
md = open(BASE + "MANUSCRIPT.md").read()
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
for s in doc.sections: s.left_margin = s.right_margin = Cm(2.2)
FIX = {"ABSTRACT": "Abstract", "1 INTRODUCTION": "1 Introduction", "2 RELATED WORK": "2 Related Work",
       "3 THE COGNITIVEGAZE DATASET": "3 The CognitiveGaze Dataset", "4 GAZE ESTIMATION MODELS": "4 Gaze Estimation Models",
       "5 DE-IDENTIFICATION METHODS": "5 De-identification Methods",
       "6 THREAT MODEL AND PRIVACY EVALUATION": "6 Threat Model and Privacy Evaluation",
       "7 IDENTITY LEAKAGE UNDER DE-IDENTIFICATION": "7 Identity Leakage under De-identification",
       "8 GAZE UTILITY UNDER DE-IDENTIFICATION": "8 Gaze Utility under De-identification",
       "9 DISCUSSION": "9 Discussion", "10 DESIGN GUIDELINES FOR RELEASING GAZE DATA": "10 Design Guidelines for Releasing Gaze Data",
       "11 LIMITATIONS": "11 Limitations", "12 CONCLUSION": "12 Conclusion", "DATA AVAILABILITY": "Data Availability",
       "SUPPLEMENTARY FIGURES": "Supplementary Figures", "SUPPLEMENTARY MATERIAL": "Supplementary Material"}
def runs(par, text):
    # **bold**, *italic*, `code`
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)", text):
        if not tok: continue
        if tok.startswith("**"): r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"): r = par.add_run(tok[1:-1]); r.font.name = "Courier New"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2: r = par.add_run(tok[1:-1]); r.italic = True
        else: par.add_run(tok)
blocks = re.split(r"\n\s*\n", md.strip())
for b in blocks:
    lines = b.split("\n")
    if b.startswith("# ") and lines[0].startswith("# Keep"):
        p = doc.add_heading(lines[0][2:], level=0); continue
    if b.startswith("# "): doc.add_heading(FIX.get(lines[0][2:], lines[0][2:]), level=1)
    if b.startswith("## "): doc.add_heading(FIX.get(lines[0][3:], lines[0][3:]), level=2)
    if b.startswith("#"):
        rest = "\n".join(lines[1:]).strip()
        if not rest: continue
        b = rest; lines = b.split("\n")
    if lines[0].startswith("|"):
        rows = [l for l in lines if l.startswith("|") and not re.match(r"^\|[-\s|:]+\|$", l)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        t = doc.add_table(rows=len(cells), cols=max(len(c) for c in cells)); t.style = "Table Grid"
        for i, row in enumerate(cells):
            for j, c in enumerate(row):
                par = t.cell(i, j).paragraphs[0]; runs(par, c)
                for r in par.runs:
                    r.font.size = Pt(8.5)
                    if i == 0: r.bold = True
        doc.add_paragraph()
        continue
    m = re.match(r"!\[(.*?)\]\((.*?)\)", b)
    if m:
        path = m.group(2)
        if path.startswith("FIGURE:"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"[{m.group(1)}: image from the previous manuscript]"); r.italic = True
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(BASE + path.replace(".pdf", ".png"), width=Cm(16))
        continue
    if lines[0].startswith("- "):
        its = []
        for l in lines:
            if l.startswith("- "): its.append(l[2:].strip())
            elif its: its[-1] += " " + l.strip()
        for it in its:
            p = doc.add_paragraph(style="List Bullet"); runs(p, it)
        continue
    text = " ".join(l.strip() for l in lines)
    p = doc.add_paragraph(); runs(p, text)
    p.paragraph_format.space_after = Pt(6)
doc.save(BASE + "MANUSCRIPT.docx")
print("saved")
