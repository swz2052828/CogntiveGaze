"""MANUSCRIPT.md -> MANUSCRIPT_highlighted.docx, with text that is new or changed
relative to an earlier MANUSCRIPT.md highlighted (word-level). Deleted text is not shown."""
import re, sys, difflib
sys.path.insert(0, "pylib")
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
BASE = "/springbrook/share/eng/esrpxk/results/paper/"
OLD = sys.argv[1]
exec(open("md2docx.py").read().split("def runs(")[0].split("md = open")[0])   # imports only
FIX = {}
src = open("md2docx.py").read()
exec(src[src.index("FIX = {"):src.index("def runs(")])

def strip(t):
    return re.sub(r"\*\*|`|(?<!\w)\*(?=\S)|(?<=\S)\*(?!\w)", "", t)

def blocks(md):
    return [b for b in re.split(r"\n\s*\n", md.strip())]

def paras(md):
    out = []
    for b in blocks(md):
        lines = b.split("\n")
        if lines[0].startswith("|"):
            out += [("row", l) for l in lines if l.startswith("|")]
        elif b.startswith("#"):
            out.append(("head", lines[0])); rest = " ".join(l.strip() for l in lines[1:]).strip()
            if rest: out.append(("par", rest))
        else:
            out.append(("par", " ".join(l.strip() for l in lines)))
    return out

old = paras(open(OLD).read())
old_par = [strip(t).split() for k, t in old if k == "par"]
old_rows = [t for k, t in old if k == "row"]
old_heads = {t for k, t in old if k == "head"}

def flags_for(text):
    w = strip(text).split()
    best, br = None, 0.0
    for o in old_par:
        sm = difflib.SequenceMatcher(None, o, w, autojunk=False)
        if sm.real_quick_ratio() < br: continue
        r = sm.ratio()
        if r > br: best, br = sm, r
    if best is None or br < 0.5:
        return [True] * len(w)
    f = [False] * len(w)
    for tag, i1, i2, j1, j2 in best.get_opcodes():
        if tag in ("insert", "replace"):
            for j in range(j1, j2): f[j] = True
    return f

def runs(par, text, flags=None, size=None, bold_all=False):
    k = 0
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)", text):
        if not tok: continue
        style = None
        if tok.startswith("**"): body, style = tok[2:-2], "b"
        elif tok.startswith("`"): body, style = tok[1:-1], "c"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2: body, style = tok[1:-1], "i"
        else: body = tok
        for piece in re.split(r"(\s+)", body):
            if not piece: continue
            r = par.add_run(piece)
            if style == "b" or bold_all: r.bold = True
            if style == "i": r.italic = True
            if style == "c": r.font.name = "Courier New"
            if size: r.font.size = Pt(size)
            if not piece.isspace():
                if flags is not None and k < len(flags) and flags[k]:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                k += 1
            elif flags is not None and 0 < k < len(flags) and flags[k - 1] and flags[k]:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW

md = open(BASE + "MANUSCRIPT.md").read()
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10.5)
for s_ in doc.sections: s_.left_margin = s_.right_margin = Cm(2.2)
note = doc.add_paragraph()
r = note.add_run("Highlighted: text new or changed since the 15 October 2026 version. Deleted text is not shown.")
r.italic = True; r.font.highlight_color = WD_COLOR_INDEX.YELLOW
n_hl = [0, 0]
for b in blocks(md):
    lines = b.split("\n")
    if b.startswith("# ") and lines[0].startswith("# Keep"):
        doc.add_heading(lines[0][2:], level=0); continue
    if b.startswith("#"):
        lvl = 1 if b.startswith("# ") else 2
        txt = lines[0][lvl + 1:]
        h = doc.add_heading(FIX.get(txt, txt), level=lvl)
        if lines[0] not in old_heads:
            for r in h.runs: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        rest = "\n".join(lines[1:]).strip()
        if not rest: continue
        b = rest; lines = b.split("\n")
    if lines[0].startswith("|"):
        rows = [l for l in lines if l.startswith("|") and not re.match(r"^\|[-\s|:]+\|$", l)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        t = doc.add_table(rows=len(cells), cols=max(len(c) for c in cells)); t.style = "Table Grid"
        for i, (raw, row) in enumerate(zip(rows, cells)):
            if raw in old_rows: match = row
            else:
                m = difflib.get_close_matches(raw, old_rows, n=1, cutoff=0.6)
                match = [c.strip() for c in m[0].strip().strip("|").split("|")] if m else None
            for j, c in enumerate(row):
                changed = match is None or j >= len(match) or match[j] != c
                fl = [changed] * len(strip(c).split())
                n_hl[1] += 1; n_hl[0] += changed
                runs(t.cell(i, j).paragraphs[0], c, fl, size=8.5, bold_all=(i == 0))
        doc.add_paragraph(); continue
    m = re.match(r"!\[(.*?)\]\((.*?)\)", b)
    if m:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if m.group(2).startswith("FIGURE:"):
            r = p.add_run(f"[{m.group(1)}: image from the previous manuscript]"); r.italic = True
        else:
            p.add_run().add_picture(BASE + m.group(2).replace(".pdf", ".png"), width=Cm(16))
        continue
    text = " ".join(l.strip() for l in lines)
    fl = flags_for(text)
    n_hl[0] += sum(fl); n_hl[1] += len(fl)
    p = doc.add_paragraph(); runs(p, text, fl); p.paragraph_format.space_after = Pt(6)
doc.save(BASE + "MANUSCRIPT_highlighted.docx")
print("saved; highlighted units", n_hl)
