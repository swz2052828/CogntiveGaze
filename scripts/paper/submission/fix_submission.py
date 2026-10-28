"""Make an anonymised, template-conformant copy of acm_submission_v1.docx (v2).
Only changes that are visible in the PDF, affect anonymity, or concern the references."""
import sys, re, copy
sys.path.insert(0, "pylib")
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/springbrook/home/e/esrpxk/.claude/uploads/e65d08c9-0d9a-465f-af83-06ea05e91b43/583ea0b2-acm_submission_v1.docx"
OUT = "/springbrook/share/eng/esrpxk/results/paper/submission/acm_submission_v2.docx"
FIG_S1 = "/springbrook/share/eng/esrpxk/results/paper/figs_core/fig_core2_utility.png"
FIG_S2 = "/springbrook/share/eng/esrpxk/results/paper/figures/fig3_threat_models.png"
d = Document(SRC)
log = []
P = d.paragraphs
def find(pred, start=0):
    for i in range(start, len(d.paragraphs)):
        if pred(d.paragraphs[i]): return d.paragraphs[i]
    raise SystemExit("not found")
def set_text(p, text, italic_markup=False):
    runs = p.runs
    proto = copy.deepcopy(runs[0]._r) if runs else None
    for r in runs: r._r.getparent().remove(r._r)
    for h in p._p.findall(qn('w:hyperlink')): p._p.remove(h)   # e-mail addresses are hyperlinks, not runs
    parts = re.split(r"(\*[^*]+\*)", text) if italic_markup else [text]
    for part in parts:
        if not part: continue
        r = p.add_run(part[1:-1] if italic_markup and part.startswith("*") else part)
        if proto is not None and proto.find(qn("w:rPr")) is not None:
            r._r.insert(0, copy.deepcopy(proto.find(qn("w:rPr"))))
        if italic_markup and part.startswith("*"): r.italic = True
def delete(p): p._p.getparent().remove(p._p)

# ---- 1. EndNote fields -> plain text (same as EndNote "Convert to Plain Text")
body = d.element.body
runs = list(body.iter(qn("w:r")))
stack, kill, n_conv, cite_groups = [], [], 0, []
for r in runs:
    fc = r.find(qn("w:fldChar"))
    it = r.find(qn("w:instrText"))
    if fc is not None and fc.get(qn("w:fldCharType")) == "begin":
        stack.append({"runs": [r], "instr": "", "phase": "instr", "result": []}); continue
    if not stack: continue
    top = stack[-1]
    if it is not None and top["phase"] == "instr":
        top["instr"] += it.text or ""; top["runs"].append(r); continue
    if fc is not None and fc.get(qn("w:fldCharType")) == "separate":
        top["runs"].append(r); top["phase"] = "result"; continue
    if fc is not None and fc.get(qn("w:fldCharType")) == "end":
        top["runs"].append(r); f = stack.pop()
        if f["instr"].strip().startswith("ADDIN EN."):
            kill += f["runs"]; n_conv += 1
            if f["instr"].strip().startswith("ADDIN EN.CITE") and f["result"]:
                cite_groups.append(f["result"])
        continue
    if top["phase"] == "instr": top["runs"].append(r)
    else: top["result"].append(r)
for r in kill:
    if r.getparent() is not None: r.getparent().remove(r)
log.append(f"Converted {n_conv} EndNote fields (citations and bibliography) to plain text")

# ---- 2. Front matter
short = find(lambda p: p.style.name == "Short Title")
title = copy.deepcopy(short._p); short._p.addprevious(title)
tp = [p for p in d.paragraphs if p._p is title][0]
tp.style = d.styles["Title_document"]; set_text(tp, "Keep Your Eyes On It")
log.append('Added the main title in the Title_document style ("Keep Your Eyes On It"); Short Title and Subtitle kept')
auth = [p for p in d.paragraphs if p.style.name in ("Authors", "Affiliation")]
set_text(auth[0], "Anonymous Author(s)"); set_text(auth[1], "Anonymous institution(s)")
for p in auth[2:]: delete(p)
log.append("Replaced the four authors, affiliations and e-mail addresses with 'Anonymous Author(s)' / 'Anonymous institution(s)'")
ref = find(lambda p: p.style.name == "ACMRef")
set_text(ref, "Anonymous Author(s). 2026. Keep Your Eyes On It: Measuring Gaze Utility and Re-identification Risk in "
              "De-identified Smartphone Gaze Data. *Proc. ACM Interact. Mob. Wearable Ubiquitous Technol.* (2026).", True)
log.append("Replaced the template's placeholder ACM Reference Format with an anonymous reference to this paper")

# ---- 3. Acknowledgments: keep only the generative-AI disclosure (required by ACM policy, not identifying)
ack_h = find(lambda p: p.text.strip().upper() == "ACKNOWLEDGMENTS")
ack_p = [p for p in d.paragraphs if p.text.startswith("The author received support")][0]
i = ack_p.text.index("During the preparation of this work")
ai_text = ack_p.text[i:]
set_text(ack_h, "Use of Generative AI"); ack_h.style = d.styles["AckHead"]
set_text(ack_p, ai_text); ack_p.style = d.styles["AckPara"]
log.append("Removed funding (UKRI Innovate UK 10031483, China Scholarship Council) and the thanks to John Taylor; "
           "kept the generative-AI disclosure under an unnumbered 'Use of Generative AI' heading")

# ---- 4. History dates (template instruction + dummy dates) removed
hd = find(lambda p: p.text.strip().upper() == "HISTORY DATES")
nxt = []
for p in d.paragraphs:
    if p._p is hd._p: start = True; continue
for p in list(d.paragraphs):
    if p.text.startswith("In case of submissions being prepared") or p.text.startswith("Received November 2019"):
        delete(p)
delete(hd)
log.append("Deleted the HISTORY DATES section (template instruction text and dummy dates)")

# ---- 5. Unnumbered back-matter headings
rh = find(lambda p: p.text.strip().upper() == "REFERENCES"); rh.style = d.styles["ReferenceHead"]
ah = find(lambda p: "APPENDICES" in p.text.upper() and p.style.name == "Head1"); ah.style = d.styles["AppendixH1"]
log.append("REFERENCES now uses ReferenceHead and the appendix heading AppendixH1, so neither is auto-numbered as a section")

# ---- 7. Three IMWUT references in the text, renumbering, references in ACM Reference Format
REFS = open("/tmp/claude-1616450260/-springbrook-share-eng-esrpxk/e65d08c9-0d9a-465f-af83-06ea05e91b43/scratchpad/ms/refs_acm.txt").read().strip().split("\n")
NEW = {
 "LAI": "H.-Y. Lai, C. G. Sodini, V. Sze, and T. Heldt. 2023. Individualized Tracking of Neurocognitive-State-Dependent Eye-Movement Features Using Mobile Devices. *Proc. ACM Interact. Mob. Wearable Ubiquitous Technol.* 7, 1, Article 19 (2023). https://doi.org/10.1145/3580843",
 "IMON": "S. Huynh, R. K. Balan, and J. Ko. 2021. iMon: Appearance-based Gaze Tracking System on Mobile Devices. *Proc. ACM Interact. Mob. Wearable Ubiquitous Technol.* 5, 4, Article 161 (2021). https://doi.org/10.1145/3494999",
 "ALH": "R. Alharbi, M. Tolba, L. C. Petito, J. Hester, and N. Alshurafa. 2019. To Mask or Not to Mask? Balancing Privacy with Visual Confirmation Utility in Activity-Oriented Wearable Cameras. *Proc. ACM Interact. Mob. Wearable Ubiquitous Technol.* 3, 3 (2019), 29 pages. https://doi.org/10.1145/3351230",
}
def insert_after(run, pieces):
    """pieces: list of (text, is_cite). Returns the citation runs created."""
    created, anchor = [], run._r
    for text, is_cite in pieces:
        nr = copy.deepcopy(run._r)
        for t in nr.findall(qn("w:t")): t.getparent().remove(t)
        t = OxmlElement("w:t"); t.text = text; t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve"); nr.append(t)
        anchor.addnext(nr); anchor = nr
        if is_cite: created.append([nr])
    return created
def insert_before_text(par_pred, marker, pieces):
    p_ = [x for x in d.paragraphs if par_pred(x)][0]
    for r in p_.runs:
        if marker in r.text:
            i = r.text.index(marker); rest = r.text[i:]; r.text = r.text[:i]
            return insert_after(r, pieces + [(rest, False)])
    raise SystemExit("marker not found: " + marker)
cite_groups += insert_before_text(lambda x: x.text.startswith("Oculomotor function is a sensitive marker"), "The data, however",
    [("Recordings from mobile devices have already been used to track eye-movement features that change with neurocognitive state ", False),
     ("[@LAI]", True), (". ", False)])
cite_groups += insert_before_text(lambda x: x.text.startswith("Smartphone gaze estimation."), "AFFNet",
    [("iMon later brought person-independent appearance-based tracking on phones to 1.49 cm on GazeCapture ", False),
     ("[@IMON]", True), (". ", False)])
fd = [x for x in d.paragraphs if x.text.startswith("Face de-identification.")][0]
cite_groups += insert_after(fd.runs[-1],
    [(" In ubiquitous computing, Alharbi et al. studied a similar privacy–utility trade-off for activity-oriented wearable cameras, "
      "examining how much image masking can protect privacy while keeping the visual confirmation that the footage is collected for ", False),
     ("[@ALH]", True), (".", False)])
log.append("Inserted three sentences citing IMWUT papers: Lai et al. 2023 (Introduction, clinical motivation), "
           "iMon (Related Work, Smartphone gaze estimation), Alharbi et al. 2019 (Related Work, Face de-identification)")

order = {id(r): i for i, r in enumerate(body.iter(qn("w:r")))}
cite_groups = [g for g in cite_groups if g and g[0].getparent() is not None]
cite_groups.sort(key=lambda g: order[id(g[0])])
def tokens(inner):
    out = []
    for tok in [x.strip() for x in inner.split(",") if x.strip()]:
        m = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", tok)
        if m: out += list(range(int(m.group(1)), int(m.group(2)) + 1))
        elif tok.startswith("@"): out.append(tok[1:])
        else: out.append(int(tok))
    return out
def gtext(g): return "".join(t.text or "" for r in g for t in r.findall(qn("w:t")))
first = []
for g in cite_groups:
    for inner in re.findall(r"\[([^\]]*)\]", gtext(g)):
        for k in tokens(inner):
            if k not in first: first.append(k)
uncited = [k for k in range(1, len(REFS) + 1) if k not in first]
assert not uncited, f"references never cited: {uncited}"
newnum = {k: i + 1 for i, k in enumerate(first)}
def fmt(nums):
    nums = sorted(set(nums)); out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ", ".join(str(n) for n in nums[i:j + 1])); i = j + 1
    return "[" + ", ".join(out) + "]"
changed = 0
for g in cite_groups:
    old = gtext(g)
    new = re.sub(r"\[([^\]]*)\]", lambda m: fmt([newnum[k] for k in tokens(m.group(1))]), old)
    changed += old != new
    ts = [t for r in g for t in r.findall(qn("w:t"))]
    ts[0].text = new
    for t in ts[1:]: t.text = ""
log.append(f"Renumbered citations by order of first appearance ({len(cite_groups)} citation groups, {changed} changed); ranges use an en dash, e.g. [33–35]")
final_refs = [NEW[k] if isinstance(k, str) else REFS[k - 1] for k in first]
bib = [x for x in d.paragraphs if x.style.name == "Bib_entry" and x.text.strip()]
while len(bib) < len(final_refs):
    clone = copy.deepcopy(bib[-1]._p); bib[-1]._p.addnext(clone)
    bib = [x for x in d.paragraphs if x.style.name == "Bib_entry" and (x.text.strip() or x._p is clone)]
for x, t in zip(bib, final_refs): set_text(x, t, True)
log.append(f"Reference list rebuilt in citation order: {len(final_refs)} entries in ACM Reference Format (49 existing + 3 IMWUT)")
renum_report = {k if isinstance(k, str) else f"old {k}": v for k, v in newnum.items()}

# ---- 8. Figures: Image style, FigureCaption style, alt text, grayscale-safe captions and plots
ALT = [
 "Circular diagram with six steps: few usable public corpora, designs converge on the same corpora, new clinical data recorded, release blocked by privacy, standard exits fail (highlighted), corpus stays unreleased, which leads back to the first step.",
 "Three panels: a photograph of a seated participant facing a large display with a smartphone and an eye tracker mounted in front of it, an illustration of the investigator at a multi-monitor control workstation, and a flow chart from data acquisition unit to processing computer to storage and analysis.",
 "Left: block diagrams of iTracker, MobileNet-V3, AFFNet and MGazeNet, each taking left-eye, right-eye and face images and a face grid, with AFFNet and MGazeNet adding a face-driven normalisation block. Right: flow chart from smartphone video through ROI extraction and the CNN to predicted gaze, with an MSE loss against EyeLink ground truth.",
 "Two source faces at the top, and two participant frames on the left with the corresponding face-swapped results for each source, in which pose, background and clothing are unchanged.",
 "Two flow diagrams starting from a participant face image. Full Synthetic passes it through the face-swap generator and uses the fully synthetic face. Hybrid uses the synthetic face for the face stream and the cropped real eye regions for the eye streams.",
 "Left: two-by-two grid of linkage ARI and verification rates for DeepPrivacy2 face and full-body modes with fixed or per-frame identity, only the full-body per-frame cell is low. Right: horizontal interval plot of eye-crop linkage ARI for every method against the eye floor.",
 "Scatter plot of linkage (vertical axis) against gaze error penalty (horizontal axis) for every release protocol. Almost all points lie near the top, the private and useful lower-left corner is empty.",
 "Grouped bar charts of gaze error penalty for each method under three train-test conditions, distinguished by hatch pattern, for Full Synthetic and Hybrid releases.",
 "Dot plot of true-accept rates for unprocessed and SimSwap eye crops under the two tests. Circles (photograph matching) lie near zero for SimSwap, triangles (within-release verification) lie between 43 and 62 percent.",
]
shapes = list(d.inline_shapes)
assert len(shapes) == 9
for sh, alt in zip(shapes, ALT):
    sh._inline.find(qn("wp:docPr")).set("descr", alt)
for p in d.paragraphs:
    if p._p.findall(".//" + qn("w:drawing")) and p.style.name != "Image":
        p.style = d.styles["Image"]
log.append("Added alt text (figure descriptions) to all 9 figures and set every image paragraph to the Image style")
caps = {"Figure 6": None, "Figure 7": None, "Figure S1": None, "Figure S2": None}
for p in d.paragraphs:
    for k in caps:
        if p.text.startswith(k + ":"): caps[k] = p
for k, p in caps.items():
    p.style = d.styles["FigureCaption"]
log.append("Captions of Figures 6, 7, S1 and S2 now use FigureCaption (centred, caption size) like Figures 1 to 5")
def rep(p, a, b):
    for r in p.runs:
        if a in r.text: r.text = r.text.replace(a, b); return True
    full = p.text
    assert a in full, (a, full[:80]); set_text(p, full.replace(a, b)); return True
rep(caps["Figure 7"], "The red dashed line marks the no-information reference", "The vertical dashed line marks the no-information reference")
rep(caps["Figure 7"], "and the grey dashed line the linkage floor", "and the horizontal dashed line the linkage floor")
rep(caps["Figure S1"], "under R→D (grey), D→R (green) and D→D (purple)", "under R→D (plain bars), D→R (hatched bars) and D→D (dotted bars)")
rep(caps["Figure S1"], "The red band marks the no-information reference and the grey band the initialisation noise floor.",
    "The horizontal band with a dashed line marks the no-information reference, and the band around zero the initialisation noise floor.")
rep(caps["Figure S2"], "Blue points show A1, matched against photographs, and orange points show verification within the release (A2).",
    "Circles show A1, matched against photographs, and triangles show verification within the release (A2).")
log.append("Figure 7, S1 and S2 captions now refer to line orientation, hatch patterns and marker shapes instead of colours")
for rid, path in (("rId22", FIG_S1), ("rId23", FIG_S2)):
    d.part.related_parts[rid]._blob = open(path, "rb").read()
log.append("Replaced Figure S1 (bars now hatched) and Figure S2 (circles vs triangles, floor label corrected) so both work in grayscale")

# ---- 8b. Author-requested wording fixes (2026-10-28)
def run_sub(pat, new, where):
    n = 0
    for p_ in d.paragraphs:
        if not where(p_): continue
        for r in p_.runs:
            if re.search(pat, r.text):
                r.text = re.sub(pat, new, r.text); n += 1
    return n
n = run_sub(r"\bGT\b", "ground truth", lambda p_: True)
log.append(f"Replaced the abbreviation GT with 'ground truth' ({n} places: Section 3 opening paragraph and Figure 3 caption)")
n = run_sub(r"refine the visual presentation of Figure X, including its layout", "refine the visual presentation, including layout",
            lambda p_: p_.style.name == "AckPara")
assert n == 1
log.append("Generative-AI statement: removed 'of Figure X' and 'its', so it no longer refers to a specific figure")
cap2 = [p_ for p_ in d.paragraphs if p_.text.startswith("Figure 2:")][0]
cap2.runs[-1].text = cap2.runs[-1].text.rstrip() + " (C) Data Flow."
log.append("Figure 2 caption: added '(C) Data Flow.'")

# ---- 9. Table header rows
for t in d.tables:
    tr = t.rows[0]._tr
    trPr = tr.find(qn("w:trPr"))
    if trPr is None:
        trPr = OxmlElement("w:trPr"); tr.insert(0, trPr)
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))
log.append(f"Marked the first row of all {len(d.tables)} tables as a header row (repeats across page breaks, read as headers by screen readers)")

# ---- 10. File properties
cp = d.core_properties
cp.author = ""; cp.last_modified_by = ""
# CCS XML, as the template instructs. python-docx caps this property at 255 characters,
# Word does not, so the dc:description element is written directly.
_desc = cp._element.get_or_add_description()
_desc.text = open("/tmp/claude-1616450260/-springbrook-share-eng-esrpxk/e65d08c9-0d9a-465f-af83-06ea05e91b43/scratchpad/ms/ccs.xml").read().strip()
log.append("Pasted the CCS 2012 XML into the Comments document property")
cp.title = "Keep Your Eyes On It: Measuring Gaze Utility and Re-identification Risk in De-identified Smartphone Gaze Data"
log.append("Cleared Author and Last Modified By in the file properties and set the document title")

# drop hyperlink relationships nobody references any more (the removed author e-mails)
xml = d.element.xml
for rId, rel in list(d.part.rels.items()):
    if rel.reltype.endswith("/hyperlink") and f'r:id="{rId}"' not in xml:
        d.part.rels.pop(rId)
log.append("Removed the orphaned e-mail hyperlink targets from the file package")
d.save(OUT)
print("\n".join(f"- {l}" for l in log))
print("saved", OUT)
